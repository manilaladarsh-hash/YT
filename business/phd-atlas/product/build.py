"""Build every customer-facing file from the verified records (schema 2.0).

Outputs (dist/):
  STEM-PhD-Atlas-2027/            the product folder that goes into the ZIP
    Atlas.html  Atlas.xlsx  Atlas.csv  Playbook.pdf  START-HERE.pdf  LICENSE.txt
  STEM-PhD-Atlas-2027-v<version>.zip
  Free-Sample-STEM-PhD-Atlas-2027.pdf

Usage: python3 product/build.py [--only explorer|sample]
Refuses to build if scripts/validate_data.py reports errors.
"""

import argparse
import base64
import collections
import csv
import gzip
import html
import json
import os
import re
import shutil
import statistics
import subprocess
import sys
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from validate_data import validate  # noqa: E402

VERSION = "2.0"
TITLE = "STEM PhD Atlas 2027"
PRODUCT = "STEM-PhD-Atlas-2027"
DIST = ROOT / "dist"
OUT = DIST / PRODUCT
DATA = ROOT / "data"
CONTENT = ROOT / "product" / "content"
TEMPLATES = ROOT / "product" / "templates"
BUILD = ROOT / ".cache" / "build"
NORTH_AMERICA = {"USA", "Canada"}
AUS_NZ = {"Australia", "New Zealand"}
CURRENCIES = ["USD", "GBP", "EUR", "AUD", "CAD", "CHF", "NZD", "SEK", "DKK", "NOK"]
PRICE_LINE = ("Priced for students: about ₹190 in India (Gumroad shows rupees and accepts UPI), $5 elsewhere. "
              "A batch licence covers up to 30 students for about ₹750. Can't afford it? Reply to your receipt for a free copy.")

SCHEMA = json.loads((DATA / "schema.json").read_text())
UF, PF = list(SCHEMA["university_fields"]), list(SCHEMA["program_fields"])
ALL_FIELDS = ["deadline", "app_fee", "fee_waiver", "gre_general", "gre_subject", "toefl_min", "ielts_min",
              "english_waiver", "supervisor", "funding", "entry_route"]
FIELD_LABELS = {
    "deadline": "Deadline (2027 entry)",
    "app_fee": "Application fee",
    "fee_waiver": "Intl fee waiver",
    "gre_general": "GRE General",
    "gre_subject": "GRE Subject test",
    "toefl_min": "TOEFL minimum",
    "ielts_min": "IELTS minimum",
    "english_waiver": "English-test waiver (English-medium Indian degree)",
    "supervisor": "Supervisor before admission",
    "funding": "PhD funding",
    "entry_route": "Entry route",
}
TAXONOMY = json.loads((DATA / "fields.json").read_text())
FIELD_BY_ID = {f["id"]: f for f in TAXONOMY["fields"]}
UNI_LIST = {u["id"]: u for u in json.loads((DATA / "universities.json").read_text())["universities"]}

FX = json.loads((DATA / "fx.json").read_text())
# Set once the Gumroad listing exists, then rebuild so the free sample links to it:
#   ATLAS_PAID_URL=https://NAME.gumroad.com/l/stem-phd-atlas-2027 python3 product/build.py
PAID_URL = os.environ.get("ATLAS_PAID_URL", "").strip()

# Short labels for universities and schools (used everywhere a name must fit).
SHORT = {
    "utoronto": "Toronto", "ubc": "UBC", "mcgill": "McGill", "mcmaster": "McMaster", "waterloo": "Waterloo",
    "ualberta": "Alberta", "calgary": "Calgary", "queens": "Queen's", "sfu": "SFU", "uottawa": "Ottawa",
    "uiuc": "UIUC", "wisconsin": "UW-Madison", "utaustin": "UT Austin", "purdue": "Purdue", "ohiostate": "Ohio State",
    "maryland": "Maryland", "stonybrook": "Stony Brook", "berkeley": "Berkeley", "uchicago": "UChicago",
    "rice": "Rice", "minnesota": "Minnesota", "stanford": "Stanford", "caltech": "Caltech", "columbia": "Columbia",
    "cornell": "Cornell", "pennstate": "Penn State", "rutgers": "Rutgers", "tamu": "Texas A&M", "ucsb": "UCSB",
    "georgiatech": "Georgia Tech", "bostonu": "Boston University",
    "cambridge": "Cambridge", "edinburgh": "Edinburgh", "epfl": "EPFL", "ethz": "ETH Zurich", "heidelberg": "Heidelberg",
    "munich-mcqst": "Munich (MCQST)", "bonn-cologne": "Bonn-Cologne (BCGS)", "sissa": "SISSA", "leiden": "Leiden",
    "bristol": "Bristol", "durham": "Durham", "imperial": "Imperial College London",
    "imprs-solar": "IMPRS Solar System (Göttingen)", "imprs-qd": "IMPRS Quantum Dynamics (Heidelberg)",
    "mps-photonics": "Max Planck School of Photonics", "mps-mattertolife": "Max Planck School Matter to Life",
    "imprs-cpqm": "IMPRS Quantum Materials (Dresden)", "imprs-hd": "IMPRS Astronomy (Heidelberg)",
    "uzh": "Zurich (UZH)", "tum": "TU Munich", "tcd": "Trinity College Dublin", "icfo": "ICFO (Barcelona)",
    "paris-saclay": "Paris-Saclay", "padova": "Padova", "ucph": "Copenhagen",
    "stockholm": "Stockholm", "kth": "KTH Stockholm", "chalmers": "Chalmers", "uppsala": "Uppsala",
    "uva": "Amsterdam (UvA)", "tudelft": "TU Delft", "aalto": "Aalto",
    "arizona": "Arizona", "brown": "Brown", "cuboulder": "CU Boulder", "duke": "Duke", "harvard": "Harvard",
    "iub": "Indiana (Bloomington)", "jhu": "Johns Hopkins", "mit": "MIT", "msu": "Michigan State", "ncsu": "NC State",
    "notredame": "Notre Dame", "pitt": "Pittsburgh", "rochester": "Rochester", "uci": "UC Irvine", "ucla": "UCLA",
    "upenn": "Penn (UPenn)", "utah": "Utah", "uwashington": "U Washington (Seattle)", "vanderbilt": "Vanderbilt",
    "virginia": "Virginia (UVA)", "yale": "Yale", "ucsd": "UC San Diego", "ucdavis": "UC Davis", "asu": "Arizona State",
    "princeton": "Princeton", "umass": "UMass Amherst", "northeastern": "Northeastern", "buffalo": "Buffalo",
    "carleton": "Carleton", "dalhousie": "Dalhousie", "guelph": "Guelph", "umanitoba": "Manitoba", "usask": "Saskatchewan",
    "western": "Western", "yorku": "York (Toronto)", "umontreal": "Montréal", "concordia": "Concordia",
    "uvic": "Victoria (Canada)", "mun": "Memorial (Newfoundland)",
    "oxford": "Oxford", "ucl": "UCL", "kcl": "King's College London", "qmul": "Queen Mary (QMUL)",
    "rhul": "Royal Holloway", "sussex": "Sussex", "southampton": "Southampton", "bath": "Bath", "surrey": "Surrey",
    "exeter": "Exeter", "glasgow": "Glasgow", "standrews": "St Andrews", "strathclyde": "Strathclyde",
    "warwick": "Warwick", "nottingham": "Nottingham", "birmingham": "Birmingham", "leeds": "Leeds",
    "sheffield": "Sheffield", "lancaster": "Lancaster", "liverpool": "Liverpool", "cardiff": "Cardiff",
    "manchester": "Manchester", "newcastle": "Newcastle", "york-uk": "York (UK)", "heriotwatt": "Heriot-Watt",
    "aberdeen": "Aberdeen", "qub": "Queen's Belfast", "loughborough": "Loughborough",
    "anu": "ANU", "unimelb": "Melbourne", "usyd": "Sydney", "unsw": "UNSW Sydney",
    "monash": "Monash", "uq": "Queensland (UQ)", "adelaide": "Adelaide", "uwa": "Western Australia (UWA)",
    "macquarie": "Macquarie", "uts": "UTS Sydney", "qut": "QUT Brisbane",
    "auckland": "Auckland", "otago": "Otago", "canterbury-nz": "Canterbury (NZ)", "vuw": "Victoria Wellington",
    "umich": "Michigan", "cmu": "Carnegie Mellon", "northwestern": "Northwestern", "nyu": "NYU", "usc": "USC",
    "ucr": "UC Riverside", "ucsc": "UC Santa Cruz", "virginiatech": "Virginia Tech", "ufl": "Florida",
    "unc": "UNC Chapel Hill", "emory": "Emory", "washu": "WashU St. Louis", "dartmouth": "Dartmouth", "rpi": "RPI",
    "njit": "NJIT", "uconn": "UConn", "tufts": "Tufts", "cwru": "Case Western", "uic": "UIC", "utdallas": "UT Dallas",
    "uhouston": "Houston", "clemson": "Clemson", "lsu": "LSU", "iastate": "Iowa State", "uiowa": "Iowa",
    "oregonstate": "Oregon State", "uoregon": "Oregon", "wsu": "Washington State", "colostate": "Colorado State",
    "mines": "Colorado School of Mines", "unm": "New Mexico", "uga": "Georgia (UGA)", "fsu": "Florida State",
    "ucf": "UCF", "temple": "Temple", "drexel": "Drexel", "syracuse": "Syracuse", "binghamton": "Binghamton",
    "utk": "Tennessee", "lmu": "LMU Munich", "kit": "KIT Karlsruhe", "rwth": "RWTH Aachen", "tuberlin": "TU Berlin",
    "fuberlin": "FU Berlin", "goettingen": "Göttingen", "freiburg": "Freiburg", "tuebingen": "Tübingen",
    "tudresden": "TU Dresden", "hamburg": "Hamburg", "stuttgart": "Stuttgart", "utrecht": "Utrecht",
    "groningen": "Groningen", "tue": "TU Eindhoven", "twente": "Twente", "wur": "Wageningen", "kuleuven": "KU Leuven",
    "ugent": "Ghent", "univie": "Vienna", "ista": "ISTA (Austria)", "tuwien": "TU Wien", "lund": "Lund", "dtu": "DTU",
    "aarhus": "Aarhus", "uio": "Oslo", "ntnu": "NTNU", "helsinki": "Helsinki", "polimi": "Politecnico di Milano",
    "polito": "Politecnico di Torino", "unibo": "Bologna", "upc": "UPC Barcelona", "unibas": "Basel", "unige": "Geneva",
    "unibe": "Bern", "sorbonne": "Sorbonne", "ipparis": "IP Paris", "uga-fr": "Grenoble Alpes", "ucd": "UCD Dublin",
}


def convert(amount, cur, to):
    rates = FX["cad_per_unit"]
    if amount is None or cur not in rates or to not in rates:
        return None
    return amount * rates[cur] / rates[to]


def fmt_inr(amount):
    """Rupees rounded to the nearest 100, with Indian digit grouping (₹1,23,400)."""
    n = int(round(amount, -2))
    head, tail = str(n)[:-3], str(n)[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return "₹" + ",".join(groups + [tail]) if groups else "₹" + tail


def area(r):
    if r["country"] in NORTH_AMERICA:
        return "North America"
    return "Australia & NZ" if r["country"] in AUS_NZ else "Europe"


def region_of(r):
    """The five buyer-facing regions used in filters and copy."""
    c = r["country"]
    if c == "USA":
        return "US"
    if c == "United Kingdom":
        return "UK"
    if c == "Canada":
        return "Canada"
    return "Australia & NZ" if c in AUS_NZ else "Europe"


# ---------------------------------------------------------------- data

def load(base=DATA):
    """Flatten university + program records into one effective record per program.

    The department's own value wins; a department value that is missing or
    'Not stated' falls back to the university-wide value. level[k] says which.
    """
    unis = {p.stem: json.loads(p.read_text()) for p in sorted((base / "universities").glob("*.json"))}
    recs = []
    for p in sorted((base / "programs").glob("*.json")):
        prog = json.loads(p.read_text())
        uni = unis[prog["university"]]
        fields, level = {}, {}
        for k in ALL_FIELDS:
            own = prog["fields"].get(k)
            if k in PF or (own and own["code"] != "Not stated"):
                fields[k], level[k] = own, "department"
            else:
                fields[k], level[k] = uni["fields"][k], "university"
        fid = prog["field"]
        recs.append({
            "id": prog["id"], "uni": prog["university"], "field": fid, "also_fields": prog.get("also_fields", []),
            "field_name": FIELD_BY_ID[fid]["name"], "group": FIELD_BY_ID[fid]["group"],
            "institution": uni["institution"], "name": prog.get("name", ""), "department": prog["department"],
            "country": uni["country"], "region": uni["region"], "program_page": prog["program_page"],
            "grad_page": uni["grad_page"], "checked": max(prog["checked"], uni["checked"]),
            "fields": fields, "level": level, "notes": prog.get("notes", ""), "uni_notes": uni.get("notes", ""),
        })
    return sorted(recs, key=lambda r: (short_uni(r), r["field_name"]))


def short_uni(r):
    return SHORT.get(r["uni"], r["institution"])


def label(r):
    """'Stanford · Chemistry', or a school's own name for physics-only schools."""
    if UNI_LIST.get(r["uni"], {}).get("kind") == "school":
        return short_uni(r)
    return f"{short_uni(r)} · {r['field_name']}"


def fee_parts(code):
    m = re.match(r"^(\d+(?:\.\d+)?) (USD|CAD|GBP|EUR|CHF|SEK|DKK|NOK|AUD|NZD)$", code)
    if m:
        return float(m.group(1)), m.group(2)
    if code.startswith("0"):
        return 0.0, ""
    return None, ""


def fee_inr(r):
    amount, cur = fee_parts(r["fields"]["app_fee"]["code"])
    if amount == 0:
        return 0
    v = convert(amount, cur, "INR") if cur else None
    return round(v, -2) if v is not None else None


def nice_date(d):
    return f"{d.strftime('%b')} {d.day}, {d.year}"


def deadline_date(r):
    code = r["fields"]["deadline"]["code"]
    return date.fromisoformat(code) if re.match(r"^\d{4}-\d{2}-\d{2}$", code) else None


def deadline_text(r):
    d = deadline_date(r)
    return nice_date(d) if d else r["fields"]["deadline"]["code"]


def flags_text(f):
    out = []
    if f.get("inferred"):
        out.append("year inferred")
    if f.get("rolling"):
        out.append("rolling")
    if f.get("conflict"):
        out.append("sources conflict")
    if f.get("time_sensitive") and f["time_sensitive"] != f.get("code"):
        out.append(f"act by {nice_date(date.fromisoformat(f['time_sensitive']))}")
    if f.get("scope"):
        out.append(f["scope"])
    return "; ".join(out)


def not_stated_share(recs):
    vals = [r["fields"][k]["code"] for r in recs for k in ALL_FIELDS]
    return sum(v == "Not stated" for v in vals) / max(len(vals), 1)


# ---------------------------------------------------------------- explorer

def explorer_payload(recs):
    """Compact JSON: URLs and university-level values are stored once."""
    urls, url_ix = [], {}

    def u(x):
        if x not in url_ix:
            url_ix[x] = len(urls)
            urls.append(x)
        return url_ix[x]

    def pack(f):
        extra = {k: f[k] for k in ("inferred", "rolling", "conflict", "time_sensitive", "scope") if f.get(k)}
        if f.get("also"):
            extra["also"] = [[a["quote"], u(a["url"])] for a in f["also"]]
        out = [f["code"], f["value"], f["quote"], u(f["url"])]
        return out + [extra] if extra else out

    unis, progs = {}, []
    for r in recs:
        if r["uni"] not in unis:
            unis[r["uni"]] = {"n": r["institution"], "s": short_uni(r), "c": r["country"], "r": r["region"],
                              "g": u(r["grad_page"]), "school": UNI_LIST.get(r["uni"], {}).get("kind") == "school",
                              "f": {}, "o": r["uni_notes"]}
        uf = unis[r["uni"]]["f"]
        own = {}
        for k in ALL_FIELDS:
            if r["level"][k] == "university":
                uf.setdefault(k, pack(r["fields"][k]))
            else:
                own[k] = pack(r["fields"][k])
        amount, cur = fee_parts(r["fields"]["app_fee"]["code"])
        usd = convert(amount, cur, "USD") if cur else (0 if amount == 0 else None)
        progs.append({"i": r["id"], "u": r["uni"], "f": r["field"], "a": r["also_fields"], "d": r["department"],
                      "n": r["name"], "p": u(r["program_page"]), "k": r["checked"], "x": own, "o": r["notes"],
                      "inr": fee_inr(r), "usd": round(usd) if usd is not None else None})
    fx = {"date": FX["date"], "source": FX["source"], "usd_inr": round(FX["cad_per_unit"]["USD"] / FX["cad_per_unit"]["INR"], 2)}
    return {"version": VERSION, "title": TITLE, "checked": max(r["checked"] for r in recs), "fx": fx,
            "taxonomy": TAXONOMY["fields"], "groups": TAXONOMY["groups"], "urls": urls, "unis": unis, "programs": progs}


def write_explorer(recs, path):
    tpl = (TEMPLATES / "atlas.html").read_text()
    raw = json.dumps(explorer_payload(recs), ensure_ascii=False, separators=(",", ":")).encode()
    packed = base64.b64encode(gzip.compress(raw, 9, mtime=0)).decode()
    marker = '"__ATLAS_GZ__"'
    if marker not in tpl:
        raise RuntimeError("explorer template marker missing")
    path.write_text(tpl.replace(marker, '"' + packed + '"', 1).replace("__TITLE__", TITLE), encoding="utf-8")
    return len(raw), path.stat().st_size


# ---------------------------------------------------------------- csv + workbook

def write_csv(recs, path):
    head = ["Program", "Field", "Also covers", "University", "Department", "Country", "Region"]
    for k in ALL_FIELDS:
        lab = FIELD_LABELS[k]
        head += [lab, f"{lab} - detail", f"{lab} - source", f"{lab} - set by"]
    head += ["Notes", "Program page", "Checked"]
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(head)
        for r in recs:
            row = [label(r), r["field_name"], ", ".join(FIELD_BY_ID[x]["name"] for x in r["also_fields"]),
                   r["institution"], r["department"], r["country"], r["region"]]
            for k in ALL_FIELDS:
                f = r["fields"][k]
                detail = f["value"] + (f" [{flags_text(f)}]" if flags_text(f) else "")
                row += [f["code"], detail, f["url"], r["level"][k]]
            row += [" ".join(x for x in (r["notes"], r["uni_notes"]) if x), r["program_page"], r["checked"]]
            w.writerow(row)


HEADER_FILL = PatternFill("solid", fgColor="1F5F8B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")


def style_header(ws, widths):
    for i, wdt in enumerate(widths, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "B2"


def write_xlsx(recs, path, values):
    wb = Workbook()
    st = wb.active
    st.title = "Start"
    lines = [
        (f"{TITLE}: planner workbook", True),
        (f"Version {VERSION} · {values['N']} programs in {values['N_FIELDS']} fields at {values['N_UNIS']} universities · data checked {values['CHECKED']}", False),
        ("", False),
        ("1. Programs: every program with field, deadline, fee, waivers, tests, supervisor rule and funding. Filter with the header arrows (start with Field and Country).", False),
        ("2. Shortlist: choose a program from the drop-down in column A; the other columns fill in automatically. Yellow cells are yours to edit.", False),
        ("3. Budget: totals your shortlisted application fees by currency, then in rupees.", False),
        ("4. Recommenders and Supervisor outreach: track requests; follow-up dates calculate themselves.", False),
        ("5. Sources: the exact official sentence and link behind every value, and whether the university or the department set it.", False),
        ("", False),
        ("'Not stated' = the official pages checked did not say (not a 'no'). Always confirm on the official page before paying or submitting.", False),
        ("Not affiliated with any university. Not immigration, visa or legal advice.", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = st.cell(row=i, column=1, value=txt)
        c.font = Font(bold=bold, size=14 if bold else 11)
    st.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Programs")
    head = ["Program", "Field", "Group", "University and department", "Country", "Region", "Deadline", "Deadline note",
            "Fee amount", "Fee currency", "Fee ≈ INR", "Intl fee waiver", "Fee waiver detail", "GRE General", "GRE Subject",
            "TOEFL minimum", "IELTS minimum", "English waiver", "English detail", "Supervisor", "Supervisor detail",
            "Funding", "Funding detail", "Entry route", "Notes", "Program page", "Checked"]
    ws.append(head)
    for r in recs:
        F = r["fields"]
        amount, cur = fee_parts(F["app_fee"]["code"])
        dl = deadline_date(r)
        inr = fee_inr(r)
        ws.append([
            label(r), r["field_name"], r["group"], f"{r['institution']} · {r['department']}", r["country"], region_of(r),
            dl if dl else F["deadline"]["code"], F["deadline"]["value"],
            amount if amount is not None else "Not stated", cur, inr if inr is not None else "",
            F["fee_waiver"]["code"], F["fee_waiver"]["value"], F["gre_general"]["code"], F["gre_subject"]["code"],
            F["toefl_min"]["code"], F["ielts_min"]["code"], F["english_waiver"]["code"], F["english_waiver"]["value"],
            F["supervisor"]["code"], F["supervisor"]["value"], F["funding"]["code"], F["funding"]["value"],
            F["entry_route"]["code"], " ".join(x for x in (r["notes"], r["uni_notes"]) if x), r["program_page"], r["checked"],
        ])
        row = ws.max_row
        if dl:
            ws.cell(row=row, column=7).number_format = "yyyy-mm-dd"
        link = ws.cell(row=row, column=26)
        link.hyperlink = r["program_page"]
        link.font = Font(color="1F5F8B", underline="single")
    style_header(ws, [34, 22, 16, 30, 10, 12, 12, 40, 10, 9, 10, 14, 40, 13, 13, 12, 11, 13, 40, 14, 40, 12, 40, 18, 40, 40, 11])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{ws.max_row}"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=False, vertical="top")
    last = ws.max_row

    sl = wb.create_sheet("Shortlist")
    sl_head = ["Program (choose)", "Country", "Deadline", "Days left", "Fee", "Currency", "Intl fee waiver",
               "GRE Subject", "English waiver", "Supervisor", "Status", "SOP ready", "Letters confirmed",
               "Fee paid / waiver", "My notes"]
    sl.append(sl_head)
    rng = f"Programs!$A$2:$A${last}"

    def look(col, r):
        return f'=IF($A{r}="","",INDEX(Programs!${col}$2:${col}${last},MATCH($A{r},{rng},0)))'

    for r in range(2, 42):
        sl[f"A{r}"].fill = INPUT_FILL
        sl[f"B{r}"] = look("E", r)
        sl[f"C{r}"] = look("G", r)
        sl[f"C{r}"].number_format = "yyyy-mm-dd"
        sl[f"D{r}"] = f'=IF(ISNUMBER(C{r}),C{r}-TODAY(),"")'
        sl[f"E{r}"] = look("I", r)
        sl[f"F{r}"] = look("J", r)
        sl[f"G{r}"] = look("L", r)
        sl[f"H{r}"] = look("O", r)
        sl[f"I{r}"] = look("R", r)
        sl[f"J{r}"] = look("T", r)
        for col in "KLMNO":
            sl[f"{col}{r}"].fill = INPUT_FILL
    dv_prog = DataValidation(type="list", formula1=rng, allow_blank=True)
    dv_status = DataValidation(type="list", formula1='"Researching,Emailed supervisor,Writing,Submitted,Interview,Offer,Declined,Accepted"', allow_blank=True)
    dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_letters = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    dv_fee = DataValidation(type="list", formula1='"Not yet,Waiver requested,Waived,Paid"', allow_blank=True)
    for dv, ref in [(dv_prog, "A2:A41"), (dv_status, "K2:K41"), (dv_yes, "L2:L41"), (dv_letters, "M2:M41"), (dv_fee, "N2:N41")]:
        sl.add_data_validation(dv)
        dv.add(ref)
    style_header(sl, [40, 9, 12, 9, 8, 9, 14, 13, 14, 14, 18, 9, 9, 16, 40])

    bd = wb.create_sheet("Budget")
    bd.column_dimensions["A"].width = 58
    cols = [get_column_letter(i) for i in range(2, 2 + len(CURRENCIES))]
    notes_col = get_column_letter(2 + len(CURRENCIES))
    for col in cols:
        bd.column_dimensions[col].width = 11
    bd.column_dimensions[notes_col].width = 40
    bd.append(["Item"] + CURRENCIES + ["Notes"])
    fee_row = ["Application fees on your shortlist (excludes rows marked Waived)"]
    for cur in CURRENCIES:
        fee_row.append(f'=SUMIFS(Shortlist!$E$2:$E$41,Shortlist!$F$2:$F$41,"{cur}",Shortlist!$N$2:$N$41,"<>Waived")')
    bd.append(fee_row + ["Calculated"])
    for item, note in [("English test (booking + extra score reports)", "Enter from the official test website"),
                       ("GRE General / Subject (booking + score reports)", "Enter only if a program on your list needs it"),
                       ("Transcripts, credential evaluation, courier", "Enter from your university / evaluator"),
                       ("Other", "")]:
        bd.append([item] + [None] * len(CURRENCIES) + [note])
    bd.append(["Total"] + [f"=SUM({c}2:{c}6)" for c in cols] + [""])
    bd.append([])
    inr = FX["cad_per_unit"]["INR"]
    bd.append([f"Rupees per 1 unit ({FX['source']}, {FX['date']}): edit to today's rate"]
              + [round(FX["cad_per_unit"][c] / inr, 2) if c in FX["cad_per_unit"] else None for c in CURRENCIES]
              + ["Pre-filled; banks and cards add a margin"])
    conv = "+".join(f'IF({c}9="",0,{c}7*{c}9)' for c in cols)
    bd.append(["Total in rupees (approximate)", f"={conv}"] + [None] * (len(CURRENCIES) - 1) + ["Currencies without a rate are left out"])
    bd["B10"].number_format = "#,##0"
    for c in bd[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in range(3, 7):
        for col in cols:
            bd[f"{col}{r}"].fill = INPUT_FILL
    for col in cols:
        bd[f"{col}9"].fill = INPUT_FILL
    bd["A7"].font = Font(bold=True)
    bd["A10"].font = Font(bold=True)

    rc = wb.create_sheet("Recommenders")
    rc.append(["Recommender", "Email", "Relationship", "Asked on", "Packet sent", "Programs invited", "Status", "Follow up by"])
    for r in range(2, 9):
        rc[f"H{r}"] = f'=IF(D{r}="","",D{r}+14)'
        rc[f"D{r}"].number_format = rc[f"H{r}"].number_format = "yyyy-mm-dd"
        for col in "ABCDEFG":
            rc[f"{col}{r}"].fill = INPUT_FILL
    dv_rs = DataValidation(type="list", formula1='"Not asked,Agreed,Invited,Submitted some,All submitted"', allow_blank=True)
    rc.add_data_validation(dv_rs)
    dv_rs.add("G2:G8")
    style_header(rc, [24, 28, 18, 12, 11, 36, 16, 12])

    so = wb.create_sheet("Supervisor outreach")
    so.append(["Professor", "University", "Group / topic", "Paper you read", "Emailed on", "Follow up on", "Reply", "Next step"])
    for r in range(2, 31):
        so[f"F{r}"] = f'=IF(E{r}="","",E{r}+10)'
        so[f"E{r}"].number_format = so[f"F{r}"].number_format = "yyyy-mm-dd"
        for col in "ABCDEGH":
            so[f"{col}{r}"].fill = INPUT_FILL
    dv_rep = DataValidation(type="list", formula1='"No reply yet,Positive,Apply and mention me,Not taking students,Other"', allow_blank=True)
    so.add_data_validation(dv_rep)
    dv_rep.add("G2:G30")
    style_header(so, [22, 22, 26, 34, 12, 12, 20, 30])

    src = wb.create_sheet("Sources")
    src.append(["Program", "Field", "Set by", "Code", "Value", "Official quote", "Source URL", "Flags", "Checked"])
    for r in recs:
        for k in ALL_FIELDS:
            f = r["fields"][k]
            quote = f["quote"] + "".join(f"  ||  {a['quote']} ({a['url']})" for a in f.get("also", []))
            src.append([label(r), FIELD_LABELS[k], r["level"][k], f["code"], f["value"], quote or "(no statement found)",
                        f["url"], flags_text(f), r["checked"]])
    style_header(src, [34, 22, 11, 14, 40, 70, 40, 18, 11])
    src.auto_filter.ref = f"A1:I{src.max_row}"
    wb.save(path)


# ---------------------------------------------------------------- documents

def node_env():
    env = dict(os.environ)
    env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def html_doc(title, body_html):
    css = (TEMPLATES / "doc.css").read_text()
    return f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{css}</style></head><body>{body_html}</body></html>"


def md_to_html(md):
    out = subprocess.run(["pandoc", "-f", "gfm+yaml_metadata_block", "-t", "html5"], input=md, capture_output=True, text=True, check=True)
    return out.stdout


def to_pdf(html_text, pdf_path, name):
    BUILD.mkdir(parents=True, exist_ok=True)
    src = BUILD / f"{name}.html"
    src.write_text(html_text, encoding="utf-8")
    subprocess.run(["node", str(ROOT / "product" / "pdf.cjs"), str(src), str(pdf_path), "--footer", TITLE], check=True, env=node_env())


def md_table(head, rows):
    esc = lambda s: str(s).replace("|", "/")
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    lines += ["| " + " | ".join(esc(c) for c in row) + " |" for row in rows]
    return "\n".join(lines)


def fill(text, values):
    def rep(m):
        key = m.group(1)
        if key not in values:
            raise KeyError(f"unfilled placeholder {key}")
        return values[key]
    out = re.sub(r"\{\{([A-Z0-9_]+)\}\}", rep, text)
    if "{{" in out:
        raise KeyError("unfilled placeholder: " + out[out.index("{{"):out.index("{{") + 40])
    return out


def names(recs, pred, limit=12, with_field=False):
    chosen = sorted({(label(r) if with_field else short_uni(r)) for r in recs if pred(r)})
    if not chosen:
        return "none in this edition"
    more = len(chosen) - limit
    return ", ".join(chosen[:limit]) + (f" and {more} more" if more > 0 else "")


def count(recs, pred):
    return str(sum(1 for r in recs if pred(r)))


def pct(recs, pred, base_pred=lambda r: True):
    base = [r for r in recs if base_pred(r)]
    return f"{round(100 * sum(1 for r in base if pred(r)) / len(base))}%" if base else "n/a"


def stats(recs):
    import playbook_stats  # product/playbook_stats.py: numbers quoted in the Playbook and sample
    return playbook_stats.compute(recs, sys.modules[__name__])


# ---------------------------------------------------------------- build

def build(only=None):
    errors, warnings, nu, npr = validate()
    if errors:
        print("\n".join(errors[:50]))
        raise SystemExit(f"data validation failed ({len(errors)} errors); not building")
    recs = load()
    labels = [label(r) for r in recs]
    dupes = {x for x in labels if labels.count(x) > 1}
    if dupes:
        raise SystemExit(f"program labels must be unique (add SHORT names): {sorted(dupes)}")
    OUT.mkdir(parents=True, exist_ok=True)
    raw, size = write_explorer(recs, OUT / "Atlas.html")
    print(f"explorer: {len(recs)} programs, data {raw / 1e6:.1f} MB raw, file {size / 1e6:.2f} MB")
    if only == "explorer":
        return None
    sys.path.insert(0, str(ROOT / "product"))
    values = stats(recs)
    import sample  # product/sample.py
    if only == "sample":
        sample.write(recs, DIST / f"Free-Sample-{PRODUCT}.pdf", values, sys.modules[__name__])
        return None
    for f in OUT.iterdir():
        if f.name != "Atlas.html":
            f.unlink()
    write_csv(recs, OUT / "Atlas.csv")
    write_xlsx(recs, OUT / "Atlas.xlsx", values)

    playbook_md = fill((CONTENT / "playbook.md").read_text(), values)
    body = md_to_html(playbook_md)
    body = body.replace("<h1", "<div class='cover'><span class='tag'>Version " + VERSION + "</span></div><h1", 1)
    to_pdf(html_doc(f"{TITLE}: Playbook", body), OUT / "Playbook.pdf", "playbook")

    start_md = fill((CONTENT / "start-here.md").read_text(), values)
    to_pdf(html_doc("Start here", md_to_html(start_md)), OUT / "START-HERE.pdf", "start-here")
    shutil.copy(CONTENT / "LICENSE.txt", OUT / "LICENSE.txt")

    sample.write(recs, DIST / f"Free-Sample-{PRODUCT}.pdf", values, sys.modules[__name__])

    zpath = DIST / f"{PRODUCT}-v{VERSION}.zip"
    if zpath.exists():
        zpath.unlink()
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(OUT.iterdir()):
            z.write(f, f"{PRODUCT}/{f.name}")
    print(f"built {zpath.relative_to(ROOT)} with {len(recs)} programs; warnings: {len(warnings)}")
    return zpath


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", choices=["explorer", "sample"])
    build(ap.parse_args().only)
