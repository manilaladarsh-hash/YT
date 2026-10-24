"""Build ops/Business-Dashboard.xlsx: the owner's weekly metrics, forecast vs
actuals, automatic validation-gate checks, expenses and study-permit hours log.

Actual results are typed in by the owner (yellow cells). Forecasts come from
research/financials.py and live on a separate sheet; the two are never mixed.

Usage: python3 ops/build_ops_workbook.py
"""

import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "research"))
import financials  # noqa: E402

HEAD = PatternFill("solid", fgColor="1F5F8B")
HFONT = Font(bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="FFF7D6")
GOOD = PatternFill("solid", fgColor="E3F2E7")
BAD = PatternFill("solid", fgColor="F8E1E1")
WEEKS = [date(2026, 10, 26) + timedelta(weeks=i) for i in range(12)]  # Mondays, Oct 26 -> Jan 11


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEAD, HFONT
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    rd = wb.active
    rd.title = "Read me"
    for i, line in enumerate([
        "STEM PhD Atlas: owner dashboard",
        "Weekly (20 min): copy last week's numbers from Gumroad into 'Weekly actuals' (yellow cells only).",
        "'Gates' turns green or red automatically against the Dec 15 decision rules.",
        "'Forecast' is the scenario model from research/financials.py ($5 list; about $2 in India via Gumroad PPP). It is NOT actual data.",
        "Free copies sent with STUDENTAID or CLUBCOPY count as neither paid orders nor revenue.",
        "'Hours' is your study-permit self-employment log. Count every business hour.",
        "Never type estimates into 'Weekly actuals'. Leave a cell blank if you don't know.",
    ], start=1):
        rd.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    rd.column_dimensions["A"].width = 110

    wa = wb.create_sheet("Weekly actuals")
    header(wa, ["Week starting", "Product views", "Sample downloads", "Paid orders", "Orders from unrelated buyers",
                "Refunds", "Gross revenue USD", "Gumroad payout USD", "Network orders (UTM whatsapp + group + alumni + club)",
                "Reddit orders (UTM reddit_profile)", "Usefulness replies", "Concrete benefits reported", "Data errors reported",
                "Orders from India", "Batch licences sold", "Free copies sent (STUDENTAID + CLUBCOPY)", "Notes"],
           [13, 12, 14, 11, 16, 9, 14, 15, 18, 12, 14, 16, 14, 12, 12, 16, 40])
    for i, wk in enumerate(WEEKS, start=2):
        wa.cell(row=i, column=1, value=wk).number_format = "yyyy-mm-dd"
        for col in range(2, 18):
            wa.cell(row=i, column=col).fill = INPUT
    tot = len(WEEKS) + 2
    wa.cell(row=tot, column=1, value="Total").font = Font(bold=True)
    for col in range(2, 17):
        letter = wa.cell(row=1, column=col).column_letter
        wa.cell(row=tot, column=col, value=f"=SUM({letter}2:{letter}{tot - 1})").font = Font(bold=True)
    chart = LineChart()
    chart.title = "Paid orders and sample downloads per week (actuals)"
    chart.y_axis.title = "count"
    chart.add_data(Reference(wa, min_col=3, min_row=1, max_row=tot - 1), titles_from_data=True)
    chart.add_data(Reference(wa, min_col=4, min_row=1, max_row=tot - 1), titles_from_data=True)
    chart.set_categories(Reference(wa, min_col=1, min_row=2, max_row=tot - 1))
    chart.height, chart.width = 8, 20
    wa.add_chart(chart, f"A{tot + 3}")

    g = wb.create_sheet("Gates")
    header(g, ["Dec 15 gate", "Target", "Actual", "Status"], [48, 10, 10, 12])
    T = f"'Weekly actuals'!"
    rows = [
        ("Paid orders", ">=10", f"={T}D{tot}", f'=IF(C2>=10,"PASS","not yet")'),
        ("Retained orders from unrelated buyers", ">=2", f"={T}E{tot}", f'=IF(C3>=2,"PASS","not yet")'),
        ("Refund rate (refunds / paid orders)", "<=10%", f'=IF({T}D{tot}=0,0,{T}F{tot}/{T}D{tot})', f'=IF(AND({T}D{tot}>0,C4<=0.1),"PASS",IF({T}D{tot}=0,"no orders","FAIL"))'),
        ("Buyers reporting a concrete benefit", ">=2", f"={T}L{tot}", f'=IF(C5>=2,"PASS","not yet")'),
        ("Critical data errors reported", "0", f"={T}M{tot}", f'=IF(C6=0,"PASS","FIX")'),
        ("Sample downloads (for the 'change' rule)", ">=50", f"={T}C{tot}", f'=IF(C7>=50,"reached","below")'),
    ]
    for r in rows:
        g.append(list(r))
    g["C4"].number_format = "0%"
    g.append([])
    g.append(["Decision (Dec 15)", "", "", '=IF(COUNTIF(D2:D6,"PASS")=5,"PROCEED",IF(OR(C7>=50,C2>=5),"CHANGE offer/price/channel, one more cycle",IF(AND(C7<20,C3=0),"STOP (keep listing passive)","CHANGE")))'])
    g["A9"].font = Font(bold=True)
    g.conditional_formatting.add("D2:D7", CellIsRule(operator="equal", formula=['"PASS"'], fill=GOOD))
    g.conditional_formatting.add("D2:D7", CellIsRule(operator="equal", formula=['"FAIL"'], fill=BAD))

    fc = wb.create_sheet("Forecast")
    header(fc, ["Scenario", "Period", "Orders", "Avg price", "Gross", "Fees", "Refunds", "Payout fee", "Net", "Owner hours", "Net per hour"],
           [14, 34, 8, 10, 10, 10, 10, 11, 10, 12, 12])
    for name, (o1, o2, ppp, disc, ref, h1, h2) in financials.SCENARIOS.items():
        for period, orders, hours in [("This season (Nov 2026-Jan 2027)", o1, h1), ("Next full season (Aug 2027-Jan 2028)", o2, h2)]:
            s = financials.season(orders, ppp, disc, ref, hours)
            fc.append([name, period, s["orders"], s["avg_price"], s["gross"], s["platform_and_card_fees"], s["refunds"],
                       s["payout_fee"], s["net"], s["owner_hours"], s["net_per_owner_hour"]])
    fc.append([])
    fc.append(["These are assumptions (see research/financials.py), not results."])

    ex = wb.create_sheet("Expenses")
    header(ex, ["Date", "Item", "Amount", "Currency", "Approved by owner?", "Notes"], [12, 40, 10, 10, 18, 40])
    ex.append([None, "Gumroad account", 0, "USD", "n/a", "No monthly fee; per-sale fees only"])
    ex.append([None, "GitHub Actions source watch", 0, "USD", "n/a", "Within free minutes"])
    for r in range(4, 20):
        for col in range(1, 7):
            ex.cell(row=r, column=col).fill = INPUT
    ex.cell(row=21, column=2, value="Total (cap this cycle: USD 30)").font = Font(bold=True)
    ex.cell(row=21, column=3, value="=SUM(C2:C20)").font = Font(bold=True)
    dv = DataValidation(type="list", formula1='"Yes,No,n/a"', allow_blank=True)
    ex.add_data_validation(dv)
    dv.add("E2:E20")

    hr = wb.create_sheet("Hours")
    header(hr, ["Date", "Task", "Minutes", "Week starting (auto)", "Notes"], [12, 44, 10, 20, 40])
    for r in range(2, 202):
        for col in (1, 2, 3, 5):
            hr.cell(row=r, column=col).fill = INPUT
        hr.cell(row=r, column=4, value=f'=IF(A{r}="","",A{r}-WEEKDAY(A{r},3))').number_format = "yyyy-mm-dd"
        hr.cell(row=r, column=1).number_format = "yyyy-mm-dd"
    hr["G1"] = "Week starting"
    hr["H1"] = "Business hours"
    hr["G1"].font = hr["H1"].font = Font(bold=True)
    for i, wk in enumerate(WEEKS, start=2):
        hr.cell(row=i, column=7, value=wk).number_format = "yyyy-mm-dd"
        hr.cell(row=i, column=8, value=f"=SUMIF($D$2:$D$201,G{i},$C$2:$C$201)/60").number_format = "0.0"
    hr["J1"] = "Reminder: self-employment counts toward the 24 h/week off-campus limit during regular terms (IRCC). Add any other off-campus work (e.g., SpecReply) here too."

    out = ROOT / "ops" / "Business-Dashboard.xlsx"
    wb.save(out)
    print("wrote", out.relative_to(ROOT))


if __name__ == "__main__":
    main()
