"""Checks for the owner-side tooling: financial model arithmetic, the
dashboard workbook's gate formulas (recalculated by LibreOffice), and the
source-watch report format.

    python3 -m pytest -q tests
"""

import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "research"))
sys.path.insert(0, str(ROOT / "scripts"))

import financials  # noqa: E402
import check_sources  # noqa: E402

DASH = ROOT / "ops" / "Business-Dashboard.xlsx"


def test_financial_model_balances():
    s = financials.season(10, 0.6, 0.0, 0.1, 10)
    gross = 10 * (0.6 * financials.PPP_PRICE + 0.4 * financials.LIST_PRICE)
    assert s["gross"] == pytest.approx(gross)
    assert s["net"] == pytest.approx(s["gross"] - s["platform_and_card_fees"] - s["refunds"] - s["payout_fee"], abs=0.02)
    assert 0 < s["net"] < s["gross"]


def test_scenarios_are_ordered():
    nets = [financials.season(o1, p, d, r, h)["net"] for (o1, _, p, d, r, h, _) in financials.SCENARIOS.values()]
    assert nets == sorted(nets)


def test_source_report_format():
    text = check_sources.report([("x", "deadline", "https://example.edu", "December 1")], {"https://b.edu": "HTTP 403"})
    assert "no longer match" in text and "| x | deadline |" in text and "could not be fetched" in text
    assert "All reachable" in check_sources.report([], {})


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_dashboard_gates_and_hours_calculate():
    wb = load_workbook(DASH)
    wa = wb["Weekly actuals"]
    # Two weeks of illustrative inputs (test only, never shipped as data).
    for col, v in zip("BCDEFGHIJKLM", [400, 30, 6, 1, 0, 60, 50, 3, 2, 4, 2, 0]):
        wa[f"{col}2"] = v
    for col, v in zip("BCDEFGHIJKLM", [300, 25, 5, 2, 1, 50, 40, 2, 2, 3, 1, 0]):
        wa[f"{col}3"] = v
    hr = wb["Hours"]
    hr["A2"], hr["B2"], hr["C2"] = date(2026, 10, 28), "Gumroad setup", 45
    hr["A3"], hr["B3"], hr["C3"] = date(2026, 10, 30), "Messages", 30
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "dash.xlsx"
        wb.save(src)
        out = Path(tmp) / "out"
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                       capture_output=True, check=True, timeout=180)
        calc = load_workbook(out / "dash.xlsx", data_only=True)
    g = calc["Gates"]
    assert g["C2"].value == 11 and g["D2"].value == "PASS"      # 11 paid orders
    assert g["C3"].value == 3 and g["D3"].value == "PASS"       # 3 unrelated
    assert g["D4"].value == "PASS"                              # 1/11 refunds = 9%
    assert g["D5"].value == "PASS" and g["D6"].value == "PASS"
    assert g["D9"].value == "PROCEED"
    h = calc["Hours"]
    assert h["H2"].value == pytest.approx(1.25)                 # 75 min in week of Oct 26
