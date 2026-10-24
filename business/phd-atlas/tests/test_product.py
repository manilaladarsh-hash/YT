"""End-to-end checks of the built kit: data integrity, every customer file,
the offline explorer in a real browser (fit engine, My plan, calendar export,
phone layout, load time), and the workbook formulas recalculated by LibreOffice.
Run after `python3 product/build.py`:

    python3 -m pytest -q tests
"""

import csv
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "product"))

import build  # noqa: E402
from validate_data import validate  # noqa: E402

OUT = build.OUT
ZIP = build.DIST / f"{build.PRODUCT}-v{build.VERSION}.zip"
SAMPLE = build.DIST / f"Free-Sample-{build.PRODUCT}.pdf"
RECS = build.load()
SIZE_BUDGET = 5_000_000


def node_env():
    env = dict(os.environ)
    env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def run_explorer(html):
    res = subprocess.run(["node", str(ROOT / "tests" / "explorer_check.cjs"), str(html)],
                         capture_output=True, text=True, env=node_env(), timeout=300)
    assert res.returncode == 0, res.stderr
    return json.loads(res.stdout)


def pdf_text(path):
    return subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True, check=True).stdout


def pdf_pages(path):
    info = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, check=True).stdout
    return int(re.search(r"Pages:\s+(\d+)", info).group(1))


def code(r, k):
    return r["fields"][k]["code"]


def expected_fit(r, today):
    """Status for the fixture applicant: bachelor's only, no GRE Subject test, English-medium degree
    instead of a test, only free or waivable fees, needs full funding. Mirrors the explorer's rules."""
    bad = warn = 0
    dl = code(r, "deadline")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", dl) and dl < today:
        bad += 1
    e = code(r, "entry_route")
    bad += e == "Master's usually required"
    warn += e == "Not stated"
    s = code(r, "gre_subject")
    bad += s == "Required"
    warn += s == "Recommended"
    if s == "Not stated" and r["country"] in ("USA", "Canada") and \
            {r["field"], *r["also_fields"]} & {"physics", "chemistry", "math", "statistics", "astronomy"}:
        warn += 1
    w = code(r, "english_waiver")
    bad += w == "No"
    warn += w in ("Conditional", "Check tool", "Not stated")
    fw = code(r, "fee_waiver")
    if not (code(r, "app_fee").startswith("0") or fw in ("No fee", "Available")):
        bad += fw in ("Domestic only", "None")
        warn += fw not in ("Domestic only", "None")
    warn += code(r, "funding") != "Guaranteed"
    return "obstacle" if bad else "check" if warn else "match"


# ---------- data ----------

def test_data_validates():
    errors, warnings, n_unis, n_programs = validate(require_cache=(ROOT / ".cache" / "pages").exists())
    assert errors == []
    assert n_programs == len(RECS) >= 100


def test_saved_page_import_rejects_bot_check_pages(tmp_path):
    bad = tmp_path / "blocked.html"
    bad.write_text("<html><body>Performing security verification</body></html>")
    res = subprocess.run([sys.executable, str(ROOT / "scripts" / "import_saved_page.py"),
                          "https://example.org/blocked", str(bad)], capture_output=True, text=True)
    assert res.returncode != 0 and "bot check" in res.stderr


def test_every_known_value_has_a_quote_and_https_source():
    for r in RECS:
        for name, f in r["fields"].items():
            assert f["url"].startswith("https://"), (r["id"], name)
            if f["code"] not in ("Not stated", "Check tool"):
                assert f["quote"].strip(), (r["id"], name)


def test_program_labels_are_unique():
    labels = [build.label(r) for r in RECS]
    assert len(labels) == len(set(labels))


# ---------- files ----------

def test_zip_contains_exactly_the_product_files():
    with zipfile.ZipFile(ZIP) as z:
        names = sorted(z.namelist())
    expected = sorted(f"{build.PRODUCT}/{n}" for n in
                      ["Atlas.csv", "Atlas.html", "Atlas.xlsx", "LICENSE.txt", "Playbook.pdf", "START-HERE.pdf"])
    assert names == expected


def test_explorer_fits_the_size_budget():
    assert (OUT / "Atlas.html").stat().st_size <= SIZE_BUDGET


def test_csv_has_one_row_per_program():
    with (OUT / "Atlas.csv").open(encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    assert len(rows) == len(RECS) + 1
    assert rows[0][:2] == ["Program", "Field"]


def test_playbook_has_all_sections_and_no_unfilled_placeholders():
    text = " ".join(pdf_text(OUT / "Playbook.pdf").split())
    for line in (build.CONTENT / "playbook.md").read_text().splitlines():
        if line.startswith("## "):
            heading = " ".join(line[3:].split())
            assert heading in text, heading
    assert "{{" not in text and "}}" not in text
    assert f"Version {build.VERSION}" in text


def test_start_here_is_one_page():
    assert pdf_pages(OUT / "START-HERE.pdf") == 1
    start = " ".join(pdf_text(OUT / "START-HERE.pdf").split())
    assert "Your first 15 minutes" in start
    # Gumroad can't email buyers until the seller's first payout, so updates must not be promised by email.
    assert "always serves the newest version" in start and "Gumroad emails you" not in start
    assert "{{" not in start


def test_sample_shows_cards_across_fields():
    text = pdf_text(SAMPLE)
    assert "FREE SAMPLE" in text and "{{" not in text
    assert pdf_pages(SAMPLE) <= 3
    import sample
    chosen = sample.pick(RECS, build)
    assert len(chosen) == sample.N_CARDS
    assert len({r["field"] for r in chosen}) >= 8
    flat = " ".join(text.split())
    for r in chosen:
        assert build.short_uni(r) in flat, build.short_uni(r)
    assert ("Get the full Atlas" in text) == bool(build.PAID_URL)


# ---------- explorer ----------

@pytest.fixture(scope="module")
def explorer():
    return run_explorer(OUT / "Atlas.html")


def test_explorer_loads_offline_without_errors(explorer):
    assert explorer["errors"] == []
    assert explorer["requests"] == [], "explorer tried to reach the network"
    assert explorer["total"] == len(RECS)
    assert explorer["firstPageCards"] == min(48, len(RECS))
    assert explorer["summaryAll"].startswith(f"{len(RECS):,} of")


def test_explorer_loads_fast(explorer):
    assert explorer["loadMs"] < 2000
    assert explorer["phoneLoadMs"] < 2000, "load time on a 4x slower CPU"


def test_explorer_search(explorer):
    also = lambda r: " ".join(build.FIELD_BY_ID[x]["name"] for x in r["also_fields"])
    hay = lambda r: " ".join(filter(None, [build.label(r), r["institution"], r["department"], r["region"], r["country"],
                                           r["field_name"], r["name"], r["notes"], r["uni_notes"], also(r)])).lower()
    hits = sum(1 for r in RECS if "waterloo" in hay(r))
    assert explorer["searchCards"] == hits >= 1


def test_fit_engine_matches_the_rules(explorer):
    today = date.today().isoformat()
    want = {r["id"]: expected_fit(r, today) for r in RECS}
    got = dict(explorer["fit"])
    wrong = [(k, got.get(k), v) for k, v in want.items() if got.get(k) != v]
    assert wrong == []
    counts = {s: sum(v == s for v in want.values()) for s in ("match", "check", "obstacle")}
    summary = " ".join(explorer["summaryFit"].split())
    assert f"{counts['match']} match" in summary and f"{counts['check']} to check" in summary
    assert f"{counts['obstacle']} obstacle" in summary
    assert explorer["feeButtonPressed"] == "true"
    if counts["match"]:
        assert "fit-match" in explorer["firstCardStatus"], "best fit sorts first"


def test_explorer_detail_shows_official_sources(explorer):
    assert explorer["dialogLinks"] and all(u.startswith("https://") for u in explorer["dialogLinks"])
    assert any(t in explorer["dialogText"].lower() for t in ("strong match", "check first", "obstacle"))


def test_my_plan_exports_and_persists(explorer):
    assert explorer["starCount"] == "3"
    assert "Add to calendar (.ics)" in explorer["planText"]
    assert explorer["csvName"] == "my-phd-plan.csv" and explorer["csvLines"] == 4  # header + 3 programs
    assert explorer["persistedStars"] == "3"
    assert explorer["planVisibleAfterReload"] is True
    assert explorer["todoPersisted"] is True


def test_calendar_file_is_valid(explorer):
    assert explorer["icsName"] == "my-phd-deadlines.ics"
    lines = explorer["ics"].split("\r\n")
    assert lines[0] == "BEGIN:VCALENDAR" and lines[-2] == "END:VCALENDAR" and lines[-1] == ""
    n = lines.count("BEGIN:VEVENT")
    assert n >= 1 and n == lines.count("END:VEVENT") == lines.count("TRIGGER:-P7D")
    assert all(len(x.encode()) <= 75 for x in lines)
    starts = [x.split(":")[1] for x in lines if x.startswith("DTSTART;VALUE=DATE:")]
    assert len(starts) == n and all(re.match(r"^\d{8}$", s) and s >= date.today().strftime("%Y%m%d") for s in starts)


def test_explorer_phone_layout(explorer):
    assert explorer["phoneOverflow"] == 0
    assert explorer["phoneGuideOverflow"] == 0


def test_customer_journey_from_downloaded_zip():
    """Unzip the delivered file elsewhere and use it the way a buyer would."""
    with tempfile.TemporaryDirectory() as tmp:
        with zipfile.ZipFile(ZIP) as z:
            z.extractall(tmp)
        folder = Path(tmp) / build.PRODUCT
        assert "Your first 15 minutes" in pdf_text(folder / "START-HERE.pdf")
        out = run_explorer(folder / "Atlas.html")
        assert out["errors"] == [] and out["requests"] == [] and out["total"] == len(RECS)
        wb = load_workbook(folder / "Atlas.xlsx")
        assert wb["Programs"].max_row == len(RECS) + 1


# ---------- workbook ----------

@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_workbook_formulas_calculate_in_libreoffice():
    """Fill two shortlist rows, recalculate with LibreOffice, read the results back."""
    wb = load_workbook(OUT / "Atlas.xlsx")
    for name in ["Start", "Programs", "Shortlist", "Budget", "Recommenders", "Supervisor outreach", "Sources"]:
        assert name in wb.sheetnames
    dated = [r for r in RECS if re.match(r"^\d{4}-", code(r, "deadline"))]
    cad = next(r for r in dated if build.fee_parts(code(r, "app_fee"))[1] == "CAD")
    usd = next(r for r in dated if build.fee_parts(code(r, "app_fee"))[1] == "USD")
    sl = wb["Shortlist"]
    sl["A2"] = build.label(cad)
    sl["A3"] = build.label(usd)
    sl["N3"] = "Waived"
    wb["Supervisor outreach"]["E2"] = date(2026, 11, 3)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "filled.xlsx"
        wb.save(src)
        outdir = Path(tmp) / "out"
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
                       capture_output=True, check=True, timeout=180)
        calc = load_workbook(outdir / "filled.xlsx", data_only=True)
    s = calc["Shortlist"]
    assert s["B2"].value == cad["country"]
    assert s["E2"].value == build.fee_parts(code(cad, "app_fee"))[0]
    assert s["F2"].value == "CAD"
    assert str(s["C2"].value).startswith(code(cad, "deadline"))
    assert isinstance(s["D2"].value, (int, float))
    assert s["H3"].value == code(usd, "gre_subject")
    b = calc["Budget"]
    col = {c.value: c.column_letter for c in b[1]}
    assert b[f"{col['CAD']}2"].value == build.fee_parts(code(cad, "app_fee"))[0]   # counted
    assert b[f"{col['USD']}2"].value == 0                                          # marked Waived, so excluded
    assert b["B10"].value > 0                                                     # converted to rupees
    assert str(calc["Supervisor outreach"]["F2"].value).startswith("2026-11-13")
